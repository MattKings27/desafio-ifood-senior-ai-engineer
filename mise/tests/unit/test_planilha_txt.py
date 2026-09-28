"""A planilha em texto: fiel à planilha, célula a célula, e sincronizada com as mudanças dela."""

from __future__ import annotations

import os
import re
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl
import pytest

from mise import planilha_txt
from mise.despensa import OrigemDoItem, carregar_despensa
from mise.dinheiro import Dinheiro
from mise.dossie import Canal, Dossie
from mise.erros import PlanilhaInvalida
from mise.mcp_server import Sessao, abrir_sessao
from mise.planilha_txt import Celula, texto_da_celula, transcrever_planilha

if TYPE_CHECKING:
    from collections.abc import Iterator

RAIZ = Path(__file__).resolve().parents[3]
PLANILHA = RAIZ / "dados" / "despensa_dona_maria.xlsx"
TRANSCRICAO = RAIZ / "dados" / "despensa_dona_maria.txt"
DOURADOS = Path(__file__).resolve().parents[1] / "golden" / "planilha_txt"

#: Todo dinheiro escrito, em qualquer formato: "R$ 9,00", "R$9", "R$ 1234.5".
DINHEIRO = re.compile(r"R\$\s*\d(?:[\d.,]*\d)?")

#: O formato que o guard-rail da conversa confere: "R$ 1.234,56", sempre com centavos.
FORMATO_DO_GUARD_RAIL = re.compile(r"R\$ \d{1,3}(?:\.\d{3})*,\d{2}")


def _fora_do_formato(texto: str) -> list[str]:
    return [v for v in DINHEIRO.findall(texto) if not FORMATO_DO_GUARD_RAIL.fullmatch(v)]


AGORA = datetime(2026, 9, 25, 13, 0, tzinfo=UTC)


def _conferir_dourado(nome: str, texto: str) -> None:
    """Compara com o arquivo dourado; `MISE_ATUALIZAR_DOURADOS=1` o regrava."""
    arquivo = DOURADOS / nome
    if os.environ.get("MISE_ATUALIZAR_DOURADOS") == "1":
        arquivo.parent.mkdir(parents=True, exist_ok=True)
        arquivo.write_text(texto, encoding="utf-8", newline="\n")
    assert texto == arquivo.read_text(encoding="utf-8")


@pytest.fixture
def sessao(tmp_path: Path) -> Iterator[Sessao]:
    s = Sessao(
        planilha=carregar_despensa(PLANILHA),
        dossie=Dossie(tmp_path / "dossie.db", relogio=lambda: AGORA),
        arquivo_txt=tmp_path / "estado" / "despensa.txt",
    )
    yield s
    s.dossie.fechar()


# --------------------------------------------------------------------------- #
# A transcrição fiel, versionada em dados/
# --------------------------------------------------------------------------- #


def test_o_arquivo_versionado_e_a_transcricao_de_hoje() -> None:
    """Se a planilha mudar, `python -m mise.planilha_txt` refaz o arquivo, e este teste lembra."""
    assert TRANSCRICAO.read_text(encoding="utf-8") == transcrever_planilha(PLANILHA), (
        "refaça com: mise/.venv/bin/python -m mise.planilha_txt "
        "dados/despensa_dona_maria.xlsx dados/despensa_dona_maria.txt"
    )


def _numero(texto: str) -> Decimal:
    return Decimal(texto.removeprefix("R$ ").replace(".", "").replace(",", "."))


def test_a_transcricao_confere_com_cada_celula_lida_pelo_openpyxl() -> None:
    livro = openpyxl.load_workbook(PLANILHA, data_only=True)
    linhas_do_texto = TRANSCRICAO.read_text(encoding="utf-8").splitlines()
    conferidas = 0
    for aba in livro.worksheets:
        titulo = next(n for n, linha in enumerate(linhas_do_texto) if f'Aba "{aba.title}"' in linha)
        for deslocamento, linha in enumerate(aba.iter_rows(), start=1):
            colunas = linhas_do_texto[titulo + deslocamento].split(" | ")
            assert len(colunas) == len(linha), (aba.title, deslocamento)
            for celula, escrita in zip(linha, colunas, strict=True):
                if isinstance(celula.value, str):
                    assert escrita == celula.value
                else:
                    assert _numero(escrita) == Decimal(str(celula.value)), celula.coordinate
                    assert escrita.startswith("R$ ") == ("R$" in celula.number_format)
                conferidas += 1
    livro.close()
    assert conferidas == 38 * 3 + 38 * 4


def test_a_transcricao_nao_tem_dinheiro_que_o_guard_rail_nao_leia() -> None:
    texto = TRANSCRICAO.read_text(encoding="utf-8")
    assert len(DINHEIRO.findall(texto)) == 37
    assert _fora_do_formato(texto) == []
    assert "R$ 663,39" not in texto  # o total é conta, e a transcrição só tem células


@pytest.mark.parametrize(
    ("celula", "texto"),
    [
        (Celula(None), ""),
        (Celula(True), "sim"),
        (Celula(False), "não"),
        (Celula(5), "5"),
        (Celula(1.5), "1,5"),
        (Celula(24.9, '"R$" #,##0.00'), "R$ 24,90"),
        (Celula(1234.5, '"R$" #,##0.00'), "R$ 1.234,50"),
        (Celula("linha\ncom | barra"), "linha com / barra"),
    ],
)
def test_texto_de_cada_celula(celula: Celula, texto: str) -> None:
    assert texto_da_celula(celula) == texto


def test_planilha_que_nao_abre_e_invalida(tmp_path: Path) -> None:
    ruim = tmp_path / "ruim.xlsx"
    ruim.write_text("não é planilha", encoding="utf-8")
    with pytest.raises(PlanilhaInvalida):
        transcrever_planilha(ruim)


def test_linhas_vazias_no_fim_da_aba_ficam_de_fora(tmp_path: Path) -> None:
    caminho = tmp_path / "curta.xlsx"
    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Outra"
    aba.append(["a", 1])
    aba.append([None, None])
    aba.append(["b", None])
    aba["A10"] = None
    livro.save(caminho)
    texto = transcrever_planilha(caminho)
    assert '== Aba "Outra" ==\na | 1\n | \nb | \n' in texto


def test_main_refaz_a_transcricao(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    saida = tmp_path / "t.txt"
    assert planilha_txt.main([str(PLANILHA), str(saida)]) == 0
    assert saida.read_text(encoding="utf-8") == TRANSCRICAO.read_text(encoding="utf-8")
    assert planilha_txt.main([str(PLANILHA)]) == 0
    assert capsys.readouterr().out == TRANSCRICAO.read_text(encoding="utf-8")
    assert planilha_txt.main([]) == 2


# --------------------------------------------------------------------------- #
# O .estado/despensa.txt
# --------------------------------------------------------------------------- #


def test_texto_sem_mudancas_e_o_dourado(sessao: Sessao) -> None:
    versao, texto = sessao.atualizar_planilha_txt()
    assert versao == 1
    assert sessao.arquivo_txt is not None
    assert sessao.arquivo_txt.read_text(encoding="utf-8") == texto
    _conferir_dourado("sem_mudancas.txt", texto)
    assert "Total: 37 itens, R$ 663,39 pagos." in texto
    assert "Gasto: R$ 0,00 · Restam: R$ 80,00" in texto
    assert _fora_do_formato(texto) == []


def test_texto_com_as_mudancas_dela_e_o_dourado(sessao: Sessao) -> None:
    sessao.mudar_despensa(
        lambda e: e.informar_embalagem("cobertura-de-chocolate", "1 kg", canal=Canal.TELA)
    )
    compra, _ = sessao.mudar_despensa(
        lambda e: e.adicionar(
            nome="Creme de leite",
            estoque=Decimal(2),
            unidade="un 200g",
            quantidade_comprada=Decimal(2),
            preco_pago=Decimal("9.00"),
            origem=OrigemDoItem.ORCAMENTO,
            canal=Canal.TELA,
        )
    )
    sessao.mudar_despensa(
        lambda e: e.adicionar(nome="Farinha de rosca", estoque=Decimal("0.5"), unidade="kg")
    )
    sessao.mudar_despensa(lambda e: e.acabou("bacon"))
    assert sessao.arquivo_txt is not None
    texto = sessao.arquivo_txt.read_text(encoding="utf-8")
    _conferir_dourado("com_mudancas.txt", texto)
    assert "Versão 6 (mudanças da senhora incluídas)" in texto
    assert "Cobertura de chocolate: Embalagem de 1 kg informada. (pela tela)" in texto
    assert "Gasto: R$ 9,00 · Restam: R$ 71,00" in texto
    assert _fora_do_formato(texto) == []

    sessao.mudar_despensa(lambda e: e.remover(compra.item_id, canal=Canal.TELA))
    texto = sessao.arquivo_txt.read_text(encoding="utf-8")
    assert "Creme de leite (despensa) · R$ 9,00 (devolvida)" in texto
    assert "Devolução: Creme de leite (despensa) · -R$ 9,00" in texto
    assert "Gasto: R$ 0,00 · Restam: R$ 80,00" in texto


def test_todo_gasto_dos_80_entra_no_extrato(sessao: Sessao) -> None:
    sessao.dossie.registrar_gasto("feira", Dinheiro.de("4.00"))
    versao, texto = sessao.planilha_txt()
    assert versao == 2
    assert "2026-09-25 10:00 · feira · R$ 4,00" in texto


def test_abrir_a_sessao_ja_grava_o_texto_ao_lado_do_dossie(tmp_path: Path) -> None:
    s = abrir_sessao(planilha=PLANILHA, banco=tmp_path / "estado" / "dossie.db")
    try:
        assert s.arquivo_txt == tmp_path / "estado" / "despensa.txt"
        assert s.arquivo_txt.read_text(encoding="utf-8").startswith("DESPENSA DA DONA MARIA")
    finally:
        s.dossie.fechar()


def test_o_caminho_do_texto_vem_do_ambiente(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    destino = tmp_path / "outro" / "planilha.txt"
    monkeypatch.setenv("MISE_DESPENSA_TXT", str(destino))
    s = abrir_sessao(planilha=PLANILHA, banco=tmp_path / "dossie.db")
    try:
        assert s.arquivo_txt == destino
        assert destino.is_file()
    finally:
        s.dossie.fechar()


def test_sem_poder_gravar_a_mudanca_continua_valendo(
    sessao: Sessao, caplog: pytest.LogCaptureFixture
) -> None:
    assert sessao.arquivo_txt is not None
    sessao.arquivo_txt.mkdir(parents=True)  # uma pasta no lugar do arquivo
    mudanca, _ = sessao.mudar_despensa(lambda e: e.acabou("bacon"))
    assert mudanca.mudou
    assert sessao.despensa["Bacon"].estoque.valor == 0
    assert "não consegui gravar" in caplog.text
    assert not list(sessao.arquivo_txt.parent.glob(".despensa.txt.*.tmp"))


def test_sem_a_planilha_original_o_texto_diz_que_nao_leu(tmp_path: Path) -> None:
    planilha = carregar_despensa(PLANILHA)
    planilha.origem = tmp_path / "sumiu.xlsx"
    s = Sessao(planilha=planilha, dossie=Dossie(tmp_path / "d.db"), arquivo_txt=None)
    try:
        versao, texto = s.atualizar_planilha_txt()
        assert versao == 1
        assert "(não consegui ler a planilha original agora)" in texto
        assert "Total: 37 itens" in texto
    finally:
        s.dossie.fechar()


def test_gravacao_atomica_nao_deixa_temporario(tmp_path: Path) -> None:
    destino = tmp_path / "a" / "despensa.txt"
    planilha_txt.escrever_atomico(destino, "um\n")
    planilha_txt.escrever_atomico(destino, "dois\n")
    assert destino.read_text(encoding="utf-8") == "dois\n"
    assert [p.name for p in destino.parent.iterdir()] == ["despensa.txt"]
    assert destino.stat().st_mode & 0o777 == 0o644
